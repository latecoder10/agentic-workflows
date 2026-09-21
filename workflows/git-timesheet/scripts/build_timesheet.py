#!/usr/bin/env python3
"""Deterministic renderer for the git-timesheet workflow.

Takes the agent's finished rows.json (one row per date+project, with the
agent's humanized description already written) and renders it to a
formatted .xlsx and a plain .csv. Also enforces the "no special characters"
and "not obviously AI-written" requirements as a mechanical safety net --
sanitizing characters is deterministic (one correct answer), so it belongs
here rather than left to the agent's discipline alone.

This script does NOT rewrite prose. It strips disallowed characters and
*flags* (via warnings in the summary it prints) rows that still look
AI-generated so the agent can revise rows.json and re-run before treating
the output as final.
"""

import argparse
import csv
import json
import re
from collections import defaultdict

ALLOWED_CHARS_RE = re.compile(r"[^A-Za-z0-9 .,'\-]")
MULTISPACE_RE = re.compile(r"\s+")

AI_TELLS = [
    "furthermore", "additionally", "leverage", "leveraged", "leveraging",
    "utilize", "utilized", "utilizing", "delve", "delved", "robust",
    "seamless", "seamlessly", "comprehensive", "in order to",
    "it is worth noting", "it's worth noting", "in conclusion",
    "moreover", "as a result of", "ensure that", "ensuring that",
]


def sanitize(text):
    if text is None:
        return ""
    cleaned = ALLOWED_CHARS_RE.sub(" ", text)
    cleaned = MULTISPACE_RE.sub(" ", cleaned).strip()
    return cleaned


def looks_ai_written(text):
    lowered = text.lower()
    return [tell for tell in AI_TELLS if tell in lowered]


def load_rows(path):
    with open(path, "r", encoding="utf-8") as f:
        return json.load(f)


def write_csv(path, header, rows):
    with open(path, "w", newline="", encoding="utf-8") as f:
        writer = csv.writer(f)
        writer.writerow(header)
        for r in rows:
            writer.writerow(r)


def write_xlsx(path, header, rows, totals_row):
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment
    from openpyxl.utils import get_column_letter

    wb = Workbook()
    ws = wb.active
    ws.title = "Timesheet"

    header_fill = PatternFill(start_color="1F4E78", end_color="1F4E78", fill_type="solid")
    header_font = Font(bold=True, color="FFFFFF")
    leave_fill = PatternFill(start_color="FCE4D6", end_color="FCE4D6", fill_type="solid")
    weekend_fill = PatternFill(start_color="EDEDED", end_color="EDEDED", fill_type="solid")

    ws.append(header)
    for cell in ws[1]:
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(vertical="center")
    ws.freeze_panes = "A2"

    status_col_idx = header.index("Status")
    for r in rows:
        ws.append(r)
        status = r[status_col_idx]
        row_idx = ws.max_row
        if status == "Leave":
            fill = leave_fill
        elif status == "Weekend":
            fill = weekend_fill
        else:
            fill = None
        if fill:
            for cell in ws[row_idx]:
                cell.fill = fill

    ws.append([])
    ws.append(totals_row)
    for cell in ws[ws.max_row]:
        cell.font = Font(bold=True)

    widths = [12, 12, 26, 10, 8, 90]
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w

    wb.save(path)


def main():
    p = argparse.ArgumentParser(description=__doc__)
    p.add_argument("--input", required=True, help="rows.json written by the agent")
    p.add_argument("--output-xlsx", required=True)
    p.add_argument("--output-csv", required=True)
    p.add_argument("--target-hours-per-day", type=float, default=8.0)
    args = p.parse_args()

    doc = load_rows(args.input)
    header = ["Date", "Day", "Project", "Status", "Hours", "Description"]
    out_rows = []
    warnings = []
    hours_by_date = defaultdict(float)

    for row in doc["rows"]:
        raw_desc = row.get("description", "")
        desc = sanitize(raw_desc)
        if desc != raw_desc.strip():
            warnings.append(f"{row['date']} ({row.get('project', '')}): stripped disallowed characters from description")
        tells = looks_ai_written(raw_desc)
        if tells:
            warnings.append(f"{row['date']} ({row.get('project', '')}): description reads as AI-generated, contains: {', '.join(tells)}")
        hours = float(row.get("hours", 0))
        hours_by_date[row["date"]] += hours
        out_rows.append([
            row["date"], row.get("weekday", ""), row.get("project", ""),
            row.get("status", ""), hours, desc,
        ])

    for date, total in sorted(hours_by_date.items()):
        matching = [r for r in doc["rows"] if r["date"] == date]
        status = matching[0].get("status") if matching else None
        if status == "Working" and abs(total - args.target_hours_per_day) > 0.01:
            warnings.append(f"{date}: total hours {total} does not match target {args.target_hours_per_day}")

    total_hours = sum(hours_by_date.values())
    working_dates = {r["date"] for r in doc["rows"] if r.get("status") == "Working"}
    leave_dates = {r["date"] for r in doc["rows"] if r.get("status") == "Leave"}
    totals_row = ["", "", "", "TOTAL", total_hours, f"{len(working_dates)} working day(s), {len(leave_dates)} leave day(s)"]

    write_csv(args.output_csv, header, out_rows)
    write_xlsx(args.output_xlsx, header, out_rows, totals_row)

    print(json.dumps({
        "xlsx": args.output_xlsx,
        "csv": args.output_csv,
        "row_count": len(out_rows),
        "total_hours": total_hours,
        "working_days": len(working_dates),
        "leave_days": len(leave_dates),
        "warnings": warnings,
    }, indent=2))


if __name__ == "__main__":
    main()
